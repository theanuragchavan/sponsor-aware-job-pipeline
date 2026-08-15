"""Regression tests for hand-edit survival across a tracker run.

The bug these lock down (found 2026-08-02): the "Apply Shortlist" sheet is
regenerated on every run and originally carried no `id` column, so anything
typed into it — status, date_applied, applied_via, notes — was silently
destroyed by the next 9am run. That sheet is the one you actually work from,
so it is where an application naturally gets logged. A month of application
logging was lost to it.

Run:  python tests/test_tracker_roundtrip.py
Exits non-zero on failure. Also collects under pytest if it's ever installed.

These build their own small workbook on purpose. A test that read the live
store would break whenever the data changed, which is how a regression test
quietly stops being run.
"""
import os
import sys
import tempfile

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import tracker  # noqa: E402


# --- fixtures ---------------------------------------------------------------
# Titles must pass ROLE_INCLUDE_TERMS and dodge SENIORITY_EXCLUDE /
# TITLE_EXCLUDE / APPLY_BLOCK_TERMS; companies must not look like agencies.
# Otherwise the row never reaches the shortlist and the test proves nothing.
def _row(jid, title, company, **over):
    row = {name: "" for name in tracker.FIELDNAMES}
    row.update({
        "id": str(jid), "title": title, "company": company,
        "location": "London", "salary_max": "65000", "tier": "1",
        "date_first_seen": "2026-08-01", "status": tracker.DEFAULT_STATUS,
        "sponsor_match": "yes", "sponsor_rating": "A",
    })
    row.update(over)
    return row


def _store(*rows):
    """Write a synthetic beautified workbook and return its path."""
    path = os.path.join(tempfile.mkdtemp(prefix="tracker-test-"), "store.xlsx")
    tracker._write_beautified_xlsx(path, list(rows))
    return path


def _sheet(path, title):
    wb = openpyxl.load_workbook(path)
    ws = wb[title]
    return wb, ws, [c.value for c in ws[1]]


def _edit(path, sheet, row_idx, **fields):
    """Type values into a sheet the way a human would, and save.

    Assigns `.value` rather than passing `value=` to `ws.cell()`: openpyxl
    treats `value=None` as "no value supplied" and skips the write, so clearing
    a cell that way silently does nothing — which once made the blank-cell test
    pass without testing anything.
    """
    wb, ws, hdr = _sheet(path, sheet)
    for key, val in fields.items():
        ws.cell(row=row_idx, column=hdr.index(key) + 1).value = val
    wb.save(path)


def _run(path):
    """One tracker run with no new jobs: load, then save."""
    tracker.save_tracker(path, tracker.load_tracker(path))


def _find(path, sheet, **match):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet]
    hdr = [c.value for c in ws[1]]
    for record in ws.iter_rows(min_row=2, values_only=True):
        row = dict(zip(hdr, record))
        if all(str(row.get(k)) == str(v) for k, v in match.items()):
            return row
    return None


# --- tests ------------------------------------------------------------------
def test_jobs_tracker_edit_survives():
    """The main sheet always round-tripped. Lock it in so a fix can't break it."""
    path = _store(_row(1, "Software Engineer", "Northwind Systems"))
    _edit(path, "Jobs Tracker", 2,
          status="applied", date_applied="2026-08-02",
          applied_via="company site", notes="kept")
    _run(path)
    row = _find(path, "Jobs Tracker", id="1")
    assert row, "row disappeared from the store"
    assert row["status"] == "applied"
    assert row["date_applied"] == "2026-08-02"
    assert row["applied_via"] == "company site"
    assert row["notes"] == "kept"


def test_shortlist_edit_is_harvested():
    """The actual bug: an application logged on the shortlist must not vanish."""
    path = _store(_row(2, "Backend Developer", "Bluebird Software"))
    _edit(path, "Apply Shortlist", 2,
          status="applied", date_applied="2026-08-02",
          applied_via="workable", notes="logged here")
    _run(path)
    row = _find(path, "Jobs Tracker", id="2")
    assert row, "row disappeared from the store"
    assert row["status"] == "applied", "shortlist edit was wiped by the run"
    assert row["date_applied"] == "2026-08-02"
    assert row["applied_via"] == "workable"
    assert row["notes"] == "logged here"


def test_harvested_row_leaves_the_shortlist():
    """Once actioned, a row is no longer a target — it drops off the view."""
    path = _store(_row(3, "Frontend Developer", "Cobalt Labs"))
    assert _find(path, "Apply Shortlist", id="3"), "row should start on the shortlist"
    _edit(path, "Apply Shortlist", 2, status="applied", date_applied="2026-08-02")
    _run(path)
    assert _find(path, "Apply Shortlist", id="3") is None, \
        "an actioned row must not stay on the shortlist"
    assert _find(path, "Jobs Tracker", id="3")["status"] == "applied"


def test_stale_shortlist_does_not_revert_a_main_sheet_edit():
    """The trap in the fix itself.

    The shortlist is generated showing "new". If the row is then marked applied
    on the main sheet, a naive harvest would read the stale "new" back over it
    and silently undo the edit at 9am.
    """
    path = _store(_row(4, "Software Engineer", "Dovetail Systems"))
    _edit(path, "Jobs Tracker", 2,
          status="applied", date_applied="2026-08-02", notes="edited on main")
    _run(path)
    row = _find(path, "Jobs Tracker", id="4")
    assert row["status"] == "applied", "stale shortlist reverted the edit"
    assert row["notes"] == "edited on main"


def test_blank_shortlist_cells_do_not_clobber():
    """The harvest only adds information; empty cells mean 'no edit'."""
    path = _store(_row(5, "Software Engineer", "Everest Data",
                       notes="keep me", applied_via="email"))
    _edit(path, "Apply Shortlist", 2, notes=None, applied_via=None)
    _run(path)
    row = _find(path, "Jobs Tracker", id="5")
    assert row["notes"] == "keep me"
    assert row["applied_via"] == "email"


def test_legacy_workbook_without_id_column_still_harvests():
    """Workbooks written before the fix have no `id` on the shortlist.

    Those edits must still survive the upgrade, matched on (company, title) —
    otherwise the fix loses exactly the data it was written to protect.
    """
    path = _store(_row(6, "Backend Developer", "Foxglove Software"))
    wb, ws, hdr = _sheet(path, "Apply Shortlist")
    ws.delete_cols(hdr.index("id") + 1)          # simulate the old format
    wb.save(path)
    _edit(path, "Apply Shortlist", 2, status="applied", notes="pre-fix edit")
    _run(path)
    row = _find(path, "Jobs Tracker", id="6")
    assert row["status"] == "applied", "legacy shortlist edit was lost"
    assert row["notes"] == "pre-fix edit"


def test_new_jobs_do_not_disturb_edited_rows():
    """A run that finds new jobs must still leave hand-edits alone."""
    path = _store(_row(7, "Software Engineer", "Gasworks Systems"))
    _edit(path, "Jobs Tracker", 2, status="applied", notes="mine")
    rows = tracker.load_tracker(path)
    tracker.merge_new_jobs(rows, [{
        "id": 999, "title": "Software Engineer",
        "company": {"display_name": "Halcyon Software"},
        "location": {"display_name": "Leeds"},
        "category": {"label": "IT Jobs"}, "created": "2026-08-02",
        "description": "", "redirect_url": "",
    }], tier=1, lookup=None)
    tracker.save_tracker(path, rows)
    assert _find(path, "Jobs Tracker", id="7")["status"] == "applied"
    assert _find(path, "Jobs Tracker", id="999"), "new job was not added"


def test_legacy_workbook_backfills_the_source_column():
    """A workbook written before `source` existed must still load usable rows.

    _load_xlsx zips header->values, so an older 20-column sheet loads with the
    new keys absent. Anything filtering on source == "adzuna" — the shortlist's
    liveness pool does — would then skip every one of the 3,000 existing rows.
    """
    path = _store(_row(1, "Software Engineer", "Northwind Systems"))

    # Rewrite it as a pre-2026-08-14 sheet: drop the two appended columns.
    wb, ws, hdr = _sheet(path, "Jobs Tracker")
    for name in ("apply_deadline", "source"):
        ws.delete_cols(hdr.index(name) + 1)
        hdr.remove(name)
    wb.save(path)

    _, _, header_now = _sheet(path, "Jobs Tracker")
    assert "source" not in header_now, "fixture did not actually go legacy"

    loaded = tracker.load_tracker(path)
    row = loaded["1"]
    assert row["source"] == tracker.SOURCE_ADZUNA, row.get("source")
    assert row["apply_deadline"] == ""


def test_namespaced_id_survives_the_roundtrip():
    """'gh:123' must not be mangled by the ".0" float-cleanup path."""
    path = _store(_row("gh:4567890", "Solutions Engineer", "Monzo"))
    loaded = tracker.load_tracker(path)
    assert "gh:4567890" in loaded, sorted(loaded)
    assert loaded["gh:4567890"]["source"] == "gh"


def test_merge_rows_is_insert_only():
    """An ATS re-fetch must not overwrite a hand-typed status."""
    existing = {"gh:1": _row("gh:1", "Solutions Engineer", "Monzo",
                             status="applied", notes="sent 2026-08-14")}
    fresh = _row("gh:1", "Solutions Engineer", "Monzo")      # status back to new
    added = tracker.merge_rows(existing, [fresh, _row("gh:2", "SWE", "Monzo")])

    assert [r["id"] for r in added] == ["gh:2"]
    assert existing["gh:1"]["status"] == "applied"
    assert existing["gh:1"]["notes"] == "sent 2026-08-14"


# --- runner -----------------------------------------------------------------
if __name__ == "__main__":
    tests = [v for k, v in sorted(globals().items()) if k.startswith("test_")]
    failed = 0
    for fn in tests:
        try:
            fn()
            print(f"PASS  {fn.__name__}")
        except AssertionError as exc:
            failed += 1
            print(f"FAIL  {fn.__name__}: {exc}")
        except Exception as exc:  # noqa: BLE001 - report, don't mask
            failed += 1
            print(f"ERROR {fn.__name__}: {type(exc).__name__}: {exc}")
    print(f"\n{len(tests) - failed}/{len(tests)} passed")
    sys.exit(1 if failed else 0)
